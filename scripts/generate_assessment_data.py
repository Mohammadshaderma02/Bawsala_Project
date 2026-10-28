import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\dana_project")
DOCS = ROOT / "src" / "backend" / "docs"
OUTPUT = ROOT / "src" / "frontend" / "dana-cotech-web" / "src" / "features" / "platform" / "assessmentData.json"


def clean(value):
    return " ".join(str(value or "").split()).strip()


interests_sheet = load_workbook(DOCS / "Career Interests.xlsx", read_only=True, data_only=True)["RIASEC"]
interests = []
for row in range(2, 74):
    interests.append({
        "id": int(interests_sheet.cell(row, 1).value),
        "en": clean(interests_sheet.cell(row, 2).value),
        "ar": clean(interests_sheet.cell(row, 3).value),
        "type": clean(interests_sheet.cell(row, 8).value),
    })

values_sheet = load_workbook(DOCS / "Values.xlsx", read_only=True, data_only=True)["Values -1.2"]
values = []
for row in range(2, 56):
    values.append({
        "id": int(values_sheet.cell(row, 1).value),
        "en": clean(values_sheet.cell(row, 2).value),
        "descriptionEn": clean(values_sheet.cell(row, 3).value),
        "ar": clean(values_sheet.cell(row, 4).value),
        "descriptionAr": clean(values_sheet.cell(row, 5).value),
    })

purpose_questions = [
    {"id": "spareTime", "en": "What do you love to do when you have spare time?", "ar": "ما الذي تحب فعله عندما يكون لديك وقت فراغ؟"},
    {"id": "enjoy", "en": "What aspects of your current job or life activities do you thoroughly enjoy?", "ar": "ما الجوانب في وظيفتك الحالية أو أنشطتك الحياتية التي تستمتع بها تماماً؟"},
    {"id": "naturalTalent", "en": "What do you naturally do well?", "ar": "ما الذي تفعله بشكل طبيعي وبإتقان؟"},
    {"id": "passion", "en": "Is there a cause, value, or quality that you feel passionate about?", "ar": "هل هناك قضية أو قيمة أو صفة تشعر بشغف تجاهها؟"},
    {"id": "successes", "en": "What do you feel have been your ten greatest successes to date?", "ar": "ما هي أعظم عشر نجاحات حققتها حتى الآن برأيك؟"},
    {"id": "lessons", "en": "What are the ten most important lessons you have learned in your life?", "ar": "ما هي أهم عشر دروس تعلمتها في حياتك؟"},
    {"id": "recurringIssues", "en": "Which issues or perceived problems have occurred repeatedly in your life?", "ar": "هل هناك قضايا أو مشاكل تكررت مراراً في حياتك؟"},
    {"id": "dreams", "en": "What do you daydream or dream about doing?", "ar": "عن ماذا تحلم أو تتخيل أن تقوم به؟"},
    {"id": "legacy", "en": "What would you want to be remembered for? What would your life be incomplete without?", "ar": "ما الذي تود أن يتم تذكرك به؟ وما الأمور التي ستكون حياتك ناقصة بدونها؟"},
    {"id": "noFailure", "en": "What would you do if you knew you could not fail?", "ar": "ماذا ستفعل إذا كنت تعلم أنك لن تفشل؟"},
    {"id": "themes", "en": "What important themes do you notice across your responses?", "ar": "ما الموضوعات الرئيسية التي تلاحظها في إجاباتك؟"},
]

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps({"interests": interests, "values": values, "purposeQuestions": purpose_questions}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Generated {len(interests)} interests and {len(values)} values at {OUTPUT}")
